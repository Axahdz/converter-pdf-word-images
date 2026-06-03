import os
import fitz  # Librería para trabajar con PDFs
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from PIL import Image

# Obtener la carpeta donde se encuentra este script
carpeta_actual = os.path.dirname(os.path.abspath(__file__))

# Buscar el primer archivo PDF disponible en la carpeta
pdfs = [f for f in os.listdir(carpeta_actual)
        if f.lower().endswith(".pdf")]

if not pdfs:
    print("No se encontró ningún PDF en la carpeta.")
    input("Presiona ENTER para salir...")
    exit()

archivo_pdf = pdfs[0]
ruta_pdf = os.path.join(carpeta_actual, archivo_pdf)

print(f"PDF encontrado: {archivo_pdf}")

# Crear una carpeta donde se guardarán temporalmente las imágenes
carpeta_imagenes = os.path.join(carpeta_actual, "imagenes_temporales")
os.makedirs(carpeta_imagenes, exist_ok=True)

# Abrir el PDF y convertir cada página en una imagen PNG
print("Convirtiendo páginas a imágenes...")

documento_pdf = fitz.open(ruta_pdf)

imagenes_generadas = []

for numero_pagina in range(len(documento_pdf)):

    pagina = documento_pdf[numero_pagina]

    # Se utiliza una escala 2x para mejorar la resolución
    pix = pagina.get_pixmap(matrix=fitz.Matrix(2, 2))

    nombre_imagen = f"pagina_{numero_pagina + 1:03}.png"

    ruta_imagen = os.path.join(
        carpeta_imagenes,
        nombre_imagen
    )

    pix.save(ruta_imagen)

    imagenes_generadas.append(ruta_imagen)

documento_pdf.close()

# Crear el documento Word de salida
print("Creando Word...")

doc = Document()

section = doc.sections[0]

# Configuración de hoja tamaño carta
section.page_width = Inches(8.5)
section.page_height = Inches(11)

# Márgenes pequeños para aprovechar mejor el espacio
section.top_margin = Inches(0.1)
section.bottom_margin = Inches(0.1)
section.left_margin = Inches(0.1)
section.right_margin = Inches(0.1)

# Calcular el área disponible dentro de la hoja
ancho_util = (
    section.page_width
    - section.left_margin
    - section.right_margin
)

alto_util = (
    section.page_height
    - section.top_margin
    - section.bottom_margin
)

# Agregar cada imagen al documento Word
for indice, ruta_img in enumerate(imagenes_generadas):

    with Image.open(ruta_img) as img:
        ancho_img, alto_img = img.size

    relacion_img = ancho_img / alto_img
    relacion_hoja = ancho_util / alto_util

    parrafo = doc.add_paragraph()
    run = parrafo.add_run()

    # Ajustar la imagen al máximo tamaño posible
    # sin deformar sus proporciones originales
    if relacion_img > relacion_hoja:
        run.add_picture(
            ruta_img,
            width=ancho_util
        )
    else:
        run.add_picture(
            ruta_img,
            height=alto_util
        )

    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Agregar salto de página excepto en la última imagen
    if indice < len(imagenes_generadas) - 1:
        run.add_break(WD_BREAK.PAGE)

# Guardar el Word utilizando el mismo nombre del PDF
nombre_word = os.path.splitext(archivo_pdf)[0] + ".docx"

ruta_word = os.path.join(
    carpeta_actual,
    nombre_word
)

doc.save(ruta_word)

print()
print("Conversión terminada")
print(f"Word generado: {nombre_word}")
print(f"Páginas procesadas: {len(imagenes_generadas)}")